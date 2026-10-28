from __future__ import annotations

from io import BytesIO

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from editable_pptx.imaging import compose, font_info, prepare_image
from editable_pptx.model import ConvertOptions, SlideSize, TextBlock


def _blank_layout(presentation: Presentation):
    for layout in presentation.slide_layouts:
        if (layout.name or "").lower() == "blank":
            return layout
    return presentation.slide_layouts[6]


def _set_typeface(run, font_name: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = etree.SubElement(r_pr, qn(tag))
        node.set("typeface", font_name)


def _add_picture(slide, image: Image.Image, placement: tuple[float, float, float, float], slide_size: SlideSize) -> None:
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    nx, ny, nw, nh = placement
    slide.shapes.add_picture(
        buffer,
        Emu(int(nx * slide_size.emu_w)),
        Emu(int(ny * slide_size.emu_h)),
        Emu(max(1, int(nw * slide_size.emu_w))),
        Emu(max(1, int(nh * slide_size.emu_h))),
    )


def _add_textbox(slide, block: TextBlock, slide_size: SlideSize, font_name: str) -> None:
    shape = slide.shapes.add_textbox(
        Emu(int(block.x * slide_size.emu_w)),
        Emu(int(block.y * slide_size.emu_h)),
        Emu(max(1, int(block.w * slide_size.emu_w))),
        Emu(max(1, int(block.h * slide_size.emu_h))),
    )
    text_frame = shape.text_frame
    text_frame.word_wrap = False
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.margin_left = Emu(0)
    text_frame.margin_right = Emu(0)
    text_frame.margin_top = Emu(0)
    text_frame.margin_bottom = Emu(0)
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.text = block.text
    run = paragraph.runs[0]
    run.font.size = Pt(block.font_pt)
    run.font.bold = False
    run.font.color.rgb = RGBColor(*block.color)
    run.font.name = font_name
    _set_typeface(run, font_name)
    try:
        shape.line.fill.background()
    except Exception:
        pass


def build_presentation(
    slides: list[tuple[Image.Image, list[TextBlock]]],
    slide_size: SlideSize,
    options: ConvertOptions,
) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Emu(slide_size.emu_w)
    presentation.slide_height = Emu(slide_size.emu_h)
    presentation.core_properties.title = "可編輯簡報"
    presentation.core_properties.author = "EditablePPTX"
    presentation.core_properties.subject = "由圖片轉換而成，文字可直接編輯"
    layout = _blank_layout(presentation)
    font_name = font_info()[2]
    for image, blocks in slides:
        prepared, placement = prepare_image(image, slide_size, options.fit)
        use_blocks = blocks if options.mode == "text" else []
        composed = compose(prepared, use_blocks, options.cover_text and options.mode == "text")
        slide = presentation.slides.add_slide(layout)
        _add_picture(slide, composed, placement, slide_size)
        if options.mode == "text":
            for block in use_blocks:
                if block.text.strip():
                    _add_textbox(slide, block, slide_size, font_name)
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
