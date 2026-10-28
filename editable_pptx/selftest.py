from __future__ import annotations

import os
import sys

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from editable_pptx.convert import analyze_image
from editable_pptx.export_pptx import build_presentation
from editable_pptx.imaging import compose, compute_slide_size, make_demo_image, prepare_image
from editable_pptx.model import ConvertOptions, TextBlock
from editable_pptx.ocr_engine import ensure_mta, join_ocr_words, recognize_lines
from editable_pptx.preview import render_preview


def main() -> None:
    if "--ocr" in sys.argv:
        run_ocr()
        return
    run_core()
    print("CORE_OK")


def run_core() -> None:
    assert join_ocr_words(["第", "三", "季"]) == "第三季"
    assert join_ocr_words(["Q3", "Report"]) == "Q3 Report"
    assert join_ocr_words(["18", "%"]) == "18%"

    image = Image.new("RGB", (400, 160), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    draw.rectangle([40, 40, 220, 100], fill=(0, 0, 0))
    block = TextBlock(
        text="第三季",
        x=0.1,
        y=0.2,
        w=0.4,
        h=0.3,
        color=(0, 0, 0),
        bg=(255, 255, 255),
        cover_box=(40, 40, 180, 60),
        font_pt=28,
    )
    covered = compose(image, [block], True)
    assert covered.getpixel((80, 70)) == (255, 255, 255)

    size = compute_slide_size("16:9", image)
    options = ConvertOptions(mode="text", lang="zh-Hant-TW", cover_text=True, show_boxes=True)
    data = build_presentation([(image, [block])], size, options)
    folder = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
    os.makedirs(folder, exist_ok=True)
    pptx_path = os.path.join(folder, "core.pptx")
    with open(pptx_path, "wb") as handle:
        handle.write(data)
    presentation = Presentation(pptx_path)
    assert presentation.slide_width == size.emu_w
    slide = presentation.slides[0]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert pictures and "第三季" in texts
    xml = slide.shapes[-1]._element.xml
    assert "Microsoft JhengHei" in xml
    assert "ea" in xml

    demo = make_demo_image()
    prepared, placement = prepare_image(demo, size, "contain")
    assert prepared.size == demo.size
    assert placement == (0.0, 0.0, 1.0, 1.0)
    preview = render_preview(demo, None, size, "contain", False, False, width=800)
    assert preview.size == (800, 450)
    assert preview.getpixel((40, 4))[0] > 140


def run_ocr() -> None:
    ensure_mta()
    from editable_pptx.ocr_engine import query_languages

    languages = query_languages()
    folder = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
    os.makedirs(folder, exist_ok=True)
    demo = make_demo_image()
    demo.save(os.path.join(folder, "demo-slide.png"), "PNG")
    lang = languages[0][0]
    lines = recognize_lines(demo, lang)
    lines_again = recognize_lines(demo, lang)
    assert lines and lines_again
    size = compute_slide_size("16:9", demo)
    options = ConvertOptions(mode="text", lang=lang, cover_text=True, show_boxes=False)
    blocks = analyze_image(demo, options, size)
    joined = "".join(block.text for block in blocks)
    report = os.path.join(folder, "report.txt")
    with open(report, "w", encoding="utf-8") as handle:
        handle.write(f"lang={lang}\n")
        for line in lines:
            handle.write(f"LINE {line.text} @ {line.x:.0f},{line.y:.0f},{line.w:.0f}x{line.h:.0f}\n")
        handle.write("--- blocks ---\n")
        for block in blocks:
            handle.write(f"{block.text} font={block.font_pt:.1f} box=({block.x:.3f},{block.y:.3f},{block.w:.3f},{block.h:.3f})\n")
    if "業務" not in joined or "回顧" not in joined:
        raise SystemExit(f"OCR missed title: {joined}")
    preview = render_preview(demo, blocks, size, "contain", True, False, width=1600)
    boxes = render_preview(demo, blocks, size, "contain", True, True, width=1600)
    preview.save(os.path.join(folder, "preview.png"), "PNG")
    boxes.save(os.path.join(folder, "preview-boxes.png"), "PNG")
    left = demo.resize((800, 450), Image.Resampling.LANCZOS)
    right = preview.resize((800, 450), Image.Resampling.LANCZOS)
    pair = Image.new("RGB", (1612, 450), (230, 226, 220))
    pair.paste(left, (0, 0))
    pair.paste(right, (812, 0))
    pair.save(os.path.join(folder, "compare.png"), "PNG")
    data = build_presentation([(demo, blocks)], size, options)
    path = os.path.join(folder, "sample.pptx")
    with open(path, "wb") as handle:
        handle.write(data)
    presentation = Presentation(path)
    texts = []
    for shape in presentation.slides[0].shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    blob = "\n".join(texts)
    if "業務" not in blob:
        raise SystemExit("pptx missing recognized text")
    with open(report, "a", encoding="utf-8") as handle:
        handle.write("--- pptx ---\n")
        handle.write(blob)
    print("OCR_OK", len(blocks))


if __name__ == "__main__":
    main()
